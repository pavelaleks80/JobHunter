"""
excel.py — output/Вакансии_ДДММГГ.xlsx.

Листы: «К отклику» (без отклика, 🆕 — новые), «Все», «Пробелы резюме», «Воронка», «Источники».
Первая колонка «Статус» — выпадающий список; отметки собираются в базу при следующем запуске.
"""
from __future__ import annotations

import re
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from ..tracker import db
from .common import salary_text, why_only

VERDICT_FILL = {"Подходит": "C6EFCE", "Частично": "FFEB9C", "Проверь вручную": "DDEBF7", "Не подходит": "F8CBAD"}
COLS = [("Статус", 14), ("Вердикт", 15), ("Соотв., %", 9), ("Балл", 7), ("Новая", 7), ("Вакансия", 46),
        ("Компания", 24), ("Зарплата", 22), ("Стоп-факторы / критичные пробелы", 34), ("Не хватает", 60),
        ("Частично", 50), ("Есть, но нет в резюме", 45), ("Совпало", 60), ("Не распознано", 40), ("Опубл.", 11),
        ("Источник", 10), ("Ключевые слова", 24), ("Почему балл", 60), ("Ссылка", 40)]
CHOICES = '"Откликнулся,Не интересно,Ответили,Приглашение,Отказ,Оффер"'


def _lines(items):
    return "\n".join(f"• {x}" for x in items)


def _row(v):
    from . import SRC_NAME
    m = v["match"]
    stop = m["stop"] + [f"критично: {h}" for h in dict.fromkeys(m["hard"])]
    return [db.STATUS_RU.get(v.get("applied"), "") or None, m["verdict"], m["fit"], v["score"],
            "🆕" if v["is_new"] else "", v["title"], v["company"], salary_text(v), "\n".join(stop), _lines(m["miss"]),
            _lines(m["partial"]), _lines(m["gap"]), _lines(m["have"]), _lines(m["unknown"]), v["published"],
            SRC_NAME.get(v["source"], v["source"]), v["tags"], v["why"], v["url"]]


def _sheet(ws, rows):
    for j, (name, width) in enumerate(COLS, 1):
        c = ws.cell(row=1, column=j, value=name)
        c.font, c.fill = Font(bold=True), PatternFill("solid", fgColor="DDEBF7")
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[c.column_letter].width = width
    ws.cell(row=1, column=1).fill = PatternFill("solid", fgColor="FFD966")
    dv = DataValidation(type="list", formula1=CHOICES, allow_blank=True)
    ws.add_data_validation(dv)
    for i, v in enumerate(rows, 2):
        for j, val in enumerate(_row(v), 1):
            c = ws.cell(row=i, column=j, value=val)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            name = COLS[j - 1][0]
            if name == "Вердикт":
                c.fill = PatternFill("solid", fgColor=VERDICT_FILL.get(val, "FFFFFF"))
            elif name == "Статус":
                c.fill = PatternFill("solid", fgColor="FFF2CC")
            elif name in ("Ссылка", "Вакансия") and v["url"]:
                c.hyperlink = v["url"]
                if name == "Ссылка":
                    c.font = Font(color="0563C1", underline="single")
        if v.get("applied"):
            for j in range(2, len(COLS) + 1):
                ws.cell(row=i, column=j).font = Font(color="999999")
        ws.row_dimensions[i].height = 90
    if rows:
        dv.add(f"A2:A{len(rows) + 1}")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions


def _gaps_sheet(ws, rows):
    miss = Counter(w for v in rows for w in set(why_only(v["match"]["miss"])))
    part = Counter(w for v in rows for w in set(why_only(v["match"]["partial"])))
    gap = Counter(w for v in rows for w in set(why_only(v["match"]["gap"])))
    ws.append([f"Сводка по {len(rows)} вакансиям: в скольких вакансиях встречается"])
    ws.append([])
    for title, cnt in (("Не хватает (нет в опыте)", miss), ("Есть у вас, но нет в резюме — дописать", gap),
                       ("Закрыто частично", part)):
        ws.append([title, "вакансий"])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        for k, n in cnt.most_common():
            ws.append([re.sub(r" \(но альтернатива.*\)$", "", k), n])
        ws.append([])
    ws.column_dimensions["A"].width = 90
    ws.column_dimensions["B"].width = 10


def write_excel(rows, status, run_dt, output_dir, funnel=None):
    from . import SRC_NAME
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / f"Вакансии_{run_dt:%d%m%y}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "К отклику"
    _sheet(ws, [v for v in rows if not v.get("applied")])
    _sheet(wb.create_sheet("Все"), rows)
    _gaps_sheet(wb.create_sheet("Пробелы резюме"), rows)
    if funnel:
        f = wb.create_sheet("Воронка")
        f.append(["Всего откликов в базе", funnel["total"]])
        for k in ("applied", "viewed", "invited", "offer", "rejected"):
            f.append([db.STATUS_RU[k], funnel["by_status"].get(k, 0)])
        f.append(["Медиана дней до ответа", funnel["median_wait"]])
        f.append([])
        f.append(["Без ответа 7+ дней — напомнить о себе или забыть"])
        for a in funnel["silent_7d"]:
            f.append([a["title"], a["company"], (a["applied_at"] or "")[:10], a["url"]])
        f.column_dimensions["A"].width = 50
        f.column_dimensions["B"].width = 30
    s = wb.create_sheet("Источники")
    s.append(["Источник", "Статус", "Подробно"])
    for k, (ok, msg) in status.items():
        s.append([SRC_NAME.get(k, k), "OK" if ok else "ОШИБКА", msg])
    s.column_dimensions["C"].width = 90
    s.append([])
    s.append(["Запуск", run_dt.strftime("%d.%m.%Y %H:%M")])
    wb.save(path)
    return path
