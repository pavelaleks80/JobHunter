"""
excel_marks.py — отметки пользователя из отчётов output/Вакансии_*.xlsx и экспорт воронки в tracker.xlsx.

В первой колонке отчёта «Статус» можно выбрать из списка: «Откликнулся», «Не интересно», «Приглашение», «Отказ»,
«Оффер» (или просто поставить «+» = откликнулся). Перед каждым запуском отметки собираются в базу.
tracker.xlsx — производный файл: пересоздаётся из базы, руками его вести не нужно.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from . import db

MARK_COL = "Статус"
MARK_VALUES = {"+": "applied", "да": "applied", "откликнулся": "applied", "не интересно": "ignored",
               "приглашение": "invited", "отказ": "rejected", "оффер": "offer", "ответили": "viewed"}


def harvest(con, output_dir: Path) -> int:
    """Собрать отметки из всех отчётов в базу. -> число новых событий."""
    added = 0
    for f in sorted(output_dir.glob("Вакансии_*.xlsx")):
        if f.name.startswith("~$"):
            continue
        try:
            wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        except Exception as e:      # noqa: BLE001 — файл открыт/битый: соберём в следующий раз
            print(f"[!!] не прочитан {f.name}: {e}")
            continue
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            hdr = list(next(rows, []) or [])
            if MARK_COL not in hdr or "Ссылка" not in hdr:
                continue
            ix = {h: i for i, h in enumerate(hdr)}
            for r in rows:
                mark, url = r[ix[MARK_COL]], r[ix["Ссылка"]]
                if mark in (None, "") or not url:
                    continue
                status = MARK_VALUES.get(str(mark).strip().lower(), "applied")
                src = str(r[ix["Источник"]]) if "Источник" in ix else ""
                if db.set_status(con, url, r[ix["Вакансия"]], r[ix["Компания"]], status, f"excel:{f.name}",
                                 source=src):
                    added += 1
        wb.close()
    return added


def export(con, path: Path) -> Path:
    """Воронка → tracker.xlsx (лист «Отклики» + «Статистика»). Файл открыт в Excel — сохраняем рядом с суффиксом."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Отклики"
    cols = [("Дата отклика", 14), ("Статус", 14), ("Вакансия", 50), ("Компания", 30), ("Источник", 10),
            ("Обновлено", 16), ("Откуда", 26), ("Ссылка", 45)]
    for j, (name, w) in enumerate(cols, 1):
        c = ws.cell(1, j, name)
        c.font, c.fill = Font(bold=True), PatternFill("solid", fgColor="DDEBF7")
        ws.column_dimensions[c.column_letter].width = w
    fills = {"invited": "C6EFCE", "offer": "A9D08E", "rejected": "F8CBAD", "viewed": "FFF2CC", "ignored": "EDEDED"}
    apps = sorted(db.applications(con), key=lambda a: a["applied_at"] or a["updated_at"] or "", reverse=True)
    for i, a in enumerate(apps, 2):
        vals = [(a["applied_at"] or "")[:10], db.STATUS_RU.get(a["status"], a["status"]), a["title"], a["company"],
                a["source"], a["updated_at"], a["origin"], a["url"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(i, j, v)
            c.alignment = Alignment(vertical="top", wrap_text=j in (3, 4))
        if a["status"] in fills:
            ws.cell(i, 2).fill = PatternFill("solid", fgColor=fills[a["status"]])
        if a["url"]:
            ws.cell(i, 8).hyperlink = a["url"]
            ws.cell(i, 8).font = Font(color="0563C1", underline="single")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    f = db.funnel(con)
    st = wb.create_sheet("Статистика")
    st.append(["Всего в воронке", f["total"]])
    for k in ("applied", "viewed", "invited", "offer", "rejected", "ignored"):
        st.append([db.STATUS_RU[k], f["by_status"].get(k, 0)])
    st.append(["Медиана дней до ответа", f["median_wait"]])
    st.append(["Без ответа 7+ дней", len(f["silent_7d"])])
    st.append([])
    st.append(["По источникам"] + [db.STATUS_RU[k] for k in ("applied", "viewed", "invited", "offer", "rejected")])
    for src, d in sorted(f["by_source"].items()):
        st.append([src] + [d.get(k, 0) for k in ("applied", "viewed", "invited", "offer", "rejected")])
    st.column_dimensions["A"].width = 28
    try:
        wb.save(path)
        return path
    except PermissionError:
        alt = path.with_name(path.stem + "_new.xlsx")
        wb.save(alt)
        return alt
