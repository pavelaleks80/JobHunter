"""
template.py — пустой шаблон трекера откликов «Мои отклики.xlsx» (создаётся командой `jobhunter init`).

Колонки шаблона совпадают с настройками tracker_import по умолчанию, поэтому для чтения достаточно
поставить в config.yaml `tracker_import: enabled: true`.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

TEMPLATE_NAME = "Мои отклики.xlsx"
SHEET = "Отклики"
# (буква, заголовок, ширина, подсказка) — порядок совпадает с ImportColumns по умолчанию
COLUMNS = [
    ("A", "Вакансия", 45, "название вакансии"),
    ("B", "Ссылка", 45, "ссылка на вакансию — по ней JobHunter узнаёт вакансию в письме"),
    ("C", "Компания", 28, ""),
    ("D", "Дата отклика", 14, "дата, например 09.10.2026"),
    ("E", "Ответили", 22, "что угодно, если работодатель ответил: дата, «звонил HR», «посмотрели»"),
    ("F", "Собеседование", 22, "что угодно, если позвали на собеседование: дата, с кем"),
    ("G", "Итог", 16, "выберите: Отказ, Снята, Оффер — или оставьте пустым"),
    ("H", "Заметки", 40, "для себя, JobHunter не читает"),
]
RESULTS = '"Отказ,Снята,Оффер"'
ROWS = 1000                                   # сколько строк заранее оформить


def create(path: Path) -> bool:
    """Создать шаблон, если файла ещё нет. -> True, если создан."""
    if path.exists():
        return False
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    head = PatternFill("solid", fgColor="DDEBF7")
    for letter, title, width, _ in COLUMNS:
        c = ws[f"{letter}1"]
        c.value, c.font, c.fill = title, Font(bold=True), head
        c.alignment = Alignment(vertical="center")
        ws.column_dimensions[letter].width = width
    for r in range(2, ROWS + 2):
        ws[f"D{r}"].number_format = "DD.MM.YYYY"
    dv = DataValidation(type="list", formula1=RESULTS, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"G2:G{ROWS + 1}")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{ROWS + 1}"

    info = wb.create_sheet("Как заполнять")
    info.column_dimensions["A"].width = 18
    info.column_dimensions["B"].width = 95
    lines = [
        ("Зачем", "Записывайте сюда отклики — JobHunter не будет присылать в письме вакансии, на которые вы уже откликнулись,"),
        ("", "и посчитает воронку: сколько ответили, позвали, отказали (jobhunter stats)."),
        ("Включить", "В workspace/config.yaml в разделе tracker_import поставьте enabled: true."),
        ("Важно", "Файл можно держать открытым в Excel, но JobHunter видит только сохранённое — после записи нажмите Ctrl+S."),
        ("", "Одна строка — один отклик, начиная со 2-й строки листа «Отклики». Названия колонок не меняйте."),
        ("", ""),
    ]
    lines += [(f"{letter} — {title}", hint) for letter, title, _, hint in COLUMNS]
    for a, b in lines:
        info.append([a, b])
    for c in info["A"]:
        c.font = Font(bold=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return True
