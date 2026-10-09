from datetime import datetime

import openpyxl

from jobhunter.config import ImportColumns, TrackerImportCfg
from jobhunter.tracker import db, xlsx_import


def _tracker(path):
    """Трекер в формате, как у автора: 3 строки шапки, B — №, C — вакансия, D — ссылка, F — компания, G — дата,
    H — дата просмотра, K — приглашение, L — итог."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Отклики"
    ws.append([None, 155, 0.59])
    ws.append([None, "№", "Вакансия", "Ссылка", "Соответствие", "Компания", "Дата отклика", "Просмотр", "HR", "", "Приглашение", "Предложение"])
    ws.append([None, "Номер", "Название", "Ссылка"])
    ws.append([None, 1, "Руководитель проектов", "https://hh.ru/vacancy/111?hhtmFrom=x", None, "Ромашка", datetime(2026, 10, 1)])
    ws.append([None, 2, "Product Owner", "https://getmatch.ru/vacancies/36556-po", None, "Лютик", datetime(2026, 10, 2),
               datetime(2026, 10, 3), None, None, None, "ОТКАЗ"])
    ws.append([None, 3, "Delivery Lead", "https://career.habr.com/vacancies/9", None, "Тест", "05.10.2026", None, None, None, "да"])
    ws.append([None, 4, "РП", "https://rabota.sber.ru/vacancy/5", None, "СБЕР", datetime(2026, 10, 6), None, None, None, None, "СНЯТО"])
    ws.append([None, 6, "Менеджер проектов", "https://hh.ru/vacancy/222", None, "МТС", datetime(2026, 10, 7), None,
               "Прошёл собес с AI-рекрутёром"])                       # контакт с HR (колонка I) — «Ответили»
    ws.append([None, 5, None, None])                                     # пустая строка — пропускается
    wb.save(path)


def _cfg(path):
    return TrackerImportCfg(enabled=True, path=str(path), sheet="Отклики", first_row=4,
                            columns=ImportColumns(title="C", url="D", company="F", date="G", viewed="H,I", invited="J,K", result="L"))


def test_read_rows(tmp_path):
    p = tmp_path / "Поиск работы.xlsx"
    _tracker(p)
    rows = xlsx_import.read_rows(_cfg(p))
    assert [r["title"] for r in rows] == ["Руководитель проектов", "Product Owner", "Delivery Lead", "РП", "Менеджер проектов"]
    assert rows[4]["statuses"] == [("viewed", None)]
    assert rows[0]["applied"] == "2026-10-01" and rows[0]["statuses"] == []
    assert rows[1]["statuses"] == [("viewed", "2026-10-03"), ("rejected", None)]
    assert rows[2]["applied"] == "2026-10-05" and rows[2]["statuses"] == [("invited", None)]
    assert rows[3]["statuses"] == [("ignored", None)]


def test_import_idempotent_and_marks_vacancies(tmp_path):
    p = tmp_path / "t.xlsx"
    _tracker(p)
    con = db.connect(tmp_path / "j.db")
    n, changed = xlsx_import.import_tracker(con, _cfg(p))
    assert n == 5 and changed > 0
    assert xlsx_import.import_tracker(con, _cfg(p))[1] == 0               # повтор — без изменений
    st = {a["url"]: (a["status"], a["source"], a["applied_at"]) for a in db.applications(con)}
    assert st["https://hh.ru/vacancy/111"] == ("applied", "hh", "2026-10-01 00:00")
    assert st["https://getmatch.ru/vacancies/36556"][:2] == ("rejected", "getmatch")
    assert st["https://career.habr.com/vacancies/9"][0] == "invited"
    assert st["https://rabota.sber.ru/vacancy/5"][0] == "ignored"
    # вакансия из письма с той же ссылкой (другие параметры) — помечена как отклик
    idx = db.applied_index(con)
    assert db.applied_status({"url": "https://getmatch.ru/vacancies/36556-po?utm=1", "title": "x", "company": "y"}, idx) == "rejected"
    # дата ответа неизвестна → в медиану не попадает; «Ответили» с датой 03.10 — попадает (1 день)
    assert db.funnel(con)["median_wait"] == 1


def test_open_file_is_read_from_memory(tmp_path, monkeypatch):
    """Файл читается целиком в память (Excel держит его открытым — прямое открытие openpyxl не нужно)."""
    p = tmp_path / "t.xlsx"
    _tracker(p)
    calls = []
    real = openpyxl.load_workbook
    monkeypatch.setattr(xlsx_import.openpyxl, "load_workbook", lambda f, **kw: calls.append(type(f).__name__) or real(f, **kw))
    xlsx_import.read_rows(_cfg(p))
    assert calls == ["BytesIO"]


def test_template_roundtrip_with_default_settings(tmp_path):
    """Шаблон «Мои отклики.xlsx» читается настройками по умолчанию — достаточно enabled: true."""
    from jobhunter.config import load_settings
    from jobhunter.tracker import template
    (tmp_path / "config.yaml").write_text("tracker_import:\n  enabled: true\n", encoding="utf-8")
    s = load_settings(tmp_path)
    p = tmp_path / template.TEMPLATE_NAME
    assert s.tracker_import.path == str(p)                            # относительный путь — от workspace
    assert template.create(p) and not template.create(p)              # второй раз не перезаписывает
    wb = openpyxl.load_workbook(p)
    ws = wb["Отклики"]
    assert [ws.cell(1, i).value for i in range(1, 9)] == ["Вакансия", "Ссылка", "Компания", "Дата отклика", "Ответили",
                                                          "Собеседование", "Итог", "Заметки"]
    assert xlsx_import.read_rows(s.tracker_import) == []              # пустой шаблон — ни одного отклика
    ws.append(["Руководитель проектов", "https://hh.ru/vacancy/1?x=1", "Ромашка", datetime(2026, 10, 9)])
    ws.append(["Product Owner", "https://getmatch.ru/vacancies/2-po", "Лютик", datetime(2026, 10, 1), "звонил HR",
               "12.10 с командой", "Отказ", "не моё"])
    wb.save(p)
    rows = xlsx_import.read_rows(s.tracker_import)
    assert [(r["title"], r["applied"]) for r in rows] == [("Руководитель проектов", "2026-10-09"), ("Product Owner", "2026-10-01")]
    assert rows[1]["statuses"] == [("viewed", None), ("invited", None), ("rejected", None)]


def test_init_creates_template(tmp_path, capsys):
    from jobhunter.cli import main
    assert main(["--workspace", str(tmp_path), "init"]) == 0
    assert (tmp_path / "Мои отклики.xlsx").exists() and (tmp_path / "config.yaml").exists()
    assert "Мои отклики.xlsx" in capsys.readouterr().out
