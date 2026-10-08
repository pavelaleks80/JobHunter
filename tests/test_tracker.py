from datetime import datetime

import openpyxl

from jobhunter.report import write_excel
from jobhunter.tracker import db, excel_marks


def _vac(**kw):
    m = {"verdict": "Подходит", "fit": 90, "stop": [], "hard": [], "miss": [], "partial": [], "gap": [], "have": [],
         "unknown": []}
    v = {"source": "habr", "id": "1", "title": "Руководитель проектов", "company": "ТестКо",
         "url": "https://career.habr.com/vacancies/1", "published": "2026-10-08", "sal_from": None, "sal_to": None,
         "currency": "RUB", "gross": None, "score": 50, "tags": "", "why": "", "is_new": True, "applied": "", "match": m}
    v.update(kw)
    return v


def test_norm_url_strips_secrets():
    assert db.norm_url("https://hh.ru/vacancy/123?loginkey=SECRET&utm=1") == "https://hh.ru/vacancy/123"
    assert db.norm_url("https://career.habr.com/vacancies/9?x=1") == "https://career.habr.com/vacancies/9"


def test_status_only_grows(tmp_path):
    con = db.connect(tmp_path / "j.db")
    url = "https://hh.ru/vacancy/5"
    assert db.set_status(con, url, "РП", "Альфа", "applied", "cli", at="2026-10-01 10:00")
    db.set_status(con, url, "РП", "Альфа", "invited", "почта", at="2026-10-04 10:00")
    db.set_status(con, url, "РП", "Альфа", "viewed", "почта", at="2026-10-05 10:00")     # не понижает
    assert db.applications(con)[0]["status"] == "invited"
    db.set_status(con, url, "РП", "Альфа", "rejected", "почта", at="2026-10-06 10:00")
    f = db.funnel(con)
    assert f["by_status"] == {"rejected": 1} and f["median_wait"] == 5


def test_applied_status_fuzzy(tmp_path):
    con = db.connect(tmp_path / "j.db")
    db.set_status(con, "", "Руководитель ИТ-проектов", "ООО Ромашка", "applied", "cli")
    idx = db.applied_index(con)
    assert db.applied_status({"url": "", "title": "Руководитель IT-проектов", "company": "Ромашка"}, idx) in ("applied", "")
    assert db.applied_status({"url": "https://x/1", "title": "Директор", "company": "Сбер"}, idx) == ""


def test_excel_roundtrip(tmp_path, settings):
    con = db.connect(tmp_path / "j.db")
    out = tmp_path / "output"
    p = write_excel([_vac(), _vac(id="2", url="https://career.habr.com/vacancies/2", title="PO")],
                    {"habr": (True, "ok")}, datetime(2026, 10, 8), out, db.funnel(con))
    wb = openpyxl.load_workbook(p)
    ws = wb["К отклику"]
    assert ws["A1"].value == "Статус"
    ws["A2"] = "+"
    ws["A3"] = "Не интересно"
    wb.save(p)
    assert excel_marks.harvest(con, out) == 2
    assert excel_marks.harvest(con, out) == 0                                  # повтор — без дублей
    st = {a["url"]: a["status"] for a in db.applications(con)}
    assert set(st.values()) == {"applied", "ignored"}
    trk = excel_marks.export(con, tmp_path / "tracker.xlsx")
    assert openpyxl.load_workbook(trk)["Отклики"].max_row == 3
